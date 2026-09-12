from __future__ import annotations

import csv
import datetime as dt
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .io import sha256_file


ROUTINGS_SHEET = "Routings"
REQUIRED_COLUMNS = (
    "Route",
    "Line",
    "Fleet",
    "Day",
    "Pairing",
    "Flight",
    "Origin",
    "Dest",
    "Dep",
    "Arr",
)
TIME_PATTERN = re.compile(r"^(?:[01]\d|2[0-3]):[0-5]\d$")


def _clean_text(value: Any, field: str, row_number: int) -> str:
    if value is None:
        raise ValueError(f"Routings row {row_number}: {field} is blank")
    text = str(value).strip()
    if not text:
        raise ValueError(f"Routings row {row_number}: {field} is blank")
    return text


def _integer(value: Any, field: str, row_number: int) -> int:
    if isinstance(value, bool):
        raise ValueError(f"Routings row {row_number}: {field} must be an integer")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"Routings row {row_number}: {field} must be an integer"
        ) from exc
    if isinstance(value, float) and value != number:
        raise ValueError(f"Routings row {row_number}: {field} must be an integer")
    return number


def _time_text(value: Any, field: str, row_number: int) -> str:
    if isinstance(value, dt.time):
        text = value.strftime("%H:%M")
    elif isinstance(value, dt.datetime):
        text = value.strftime("%H:%M")
    elif isinstance(value, (int, float)):
        minutes = round((float(value) % 1) * 24 * 60) % (24 * 60)
        text = f"{minutes // 60:02d}:{minutes % 60:02d}"
    else:
        text = _clean_text(value, field, row_number)
        if re.fullmatch(r"\d:\d{2}", text):
            text = f"0{text}"
    if not TIME_PATTERN.fullmatch(text):
        raise ValueError(
            f"Routings row {row_number}: {field} must be a 24-hour HH:MM value; got {value!r}"
        )
    return text


def load_city_information(path: Path) -> list[dict[str, Any]]:
    cities: list[dict[str, Any]] = []
    seen: set[str] = set()
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        required = {"Code", "City", "Status", "Active"}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"City information is missing columns: {sorted(missing)}")
        for row_number, row in enumerate(reader, start=2):
            code = (row.get("Code") or "").strip().upper()
            if not code:
                raise ValueError(f"City information row {row_number}: Code is blank")
            if code in seen:
                raise ValueError(f"City information contains duplicate Code {code}")
            seen.add(code)
            status = (row.get("Status") or "").strip()
            role = {
                "Hub": "hub",
                "Focus City": "focus_city",
                "Destination": "destination",
            }.get(status, "unknown")
            hub_assignments = [
                item.strip()
                for item in (row.get("Hub_Assignment") or "").split("/")
                if item.strip() and item.strip() != "-"
            ]
            cities.append(
                {
                    "code": code,
                    "name": (row.get("City") or "").strip(),
                    "displayName": (
                        (row.get("Long_City") or "").strip()
                        or (row.get("City") or "").strip()
                    ),
                    "sourceOrder": len(cities),
                    "group": (row.get("Group") or "").strip() or None,
                    "role": role,
                    "active": (row.get("Active") or "").strip().upper() == "Y",
                    "timezone": (row.get("Timezone") or "").strip() or None,
                    "hubAssignments": hub_assignments,
                    "maintenanceBase": (row.get("MX_Base") or "").strip().upper() == "Y",
                    "gateAllocationOverride": _optional_int(
                        row.get("Gate_Allocation_Override"), code, "Gate_Allocation_Override"
                    ),
                    "standAllocationOverride": _optional_int(
                        row.get("Stand_Allocation_Override"), code, "Stand_Allocation_Override"
                    ),
                    "latitude": _optional_float(row.get("Latitude"), code, "Latitude"),
                    "longitude": _optional_float(row.get("Longitude"), code, "Longitude"),
                    "runwayLengthFeet": _optional_int(
                        row.get("Runway_Length_Ft"), code, "Runway_Length_Ft"
                    ),
                    "elevationFeet": _optional_int(
                        row.get("Elevation_Ft"), code, "Elevation_Ft"
                    ),
                }
            )
    return sorted(cities, key=lambda city: city["code"])


def _minute_of_day(value: str) -> int:
    hours, minutes = (int(part) for part in value.split(":"))
    return hours * 60 + minutes


def _optional_int(value: Any, code: str, field: str) -> int | None:
    text = "" if value is None else str(value).strip()
    if not text:
        return None
    try:
        number = float(text)
    except ValueError as exc:
        raise ValueError(f"{code}: {field} must be numeric") from exc
    if not number.is_integer():
        raise ValueError(f"{code}: {field} must be a whole number")
    return int(number)


def _optional_float(value: Any, code: str, field: str) -> float | None:
    text = "" if value is None else str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"{code}: {field} must be numeric") from exc


def load_routings(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if ROUTINGS_SHEET not in workbook.sheetnames:
        raise ValueError(f"Workbook does not contain required sheet {ROUTINGS_SHEET!r}")
    sheet = workbook[ROUTINGS_SHEET]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if headers is None:
        raise ValueError("Routings sheet is empty")
    header_index = {str(value).strip(): index for index, value in enumerate(headers) if value is not None}
    missing = [column for column in REQUIRED_COLUMNS if column not in header_index]
    if missing:
        raise ValueError(f"Routings sheet is missing columns: {missing}")

    route_sequences: defaultdict[int, int] = defaultdict(int)
    legs: list[dict[str, Any]] = []
    for row_number, values in enumerate(rows, start=2):
        if all(value is None for value in values):
            continue
        value = lambda column: values[header_index[column]]
        route = _integer(value("Route"), "Route", row_number)
        route_sequences[route] += 1
        flight = _integer(value("Flight"), "Flight", row_number)
        departure = _time_text(value("Dep"), "Dep", row_number)
        arrival = _time_text(value("Arr"), "Arr", row_number)
        departure_minute = _minute_of_day(departure)
        arrival_minute = _minute_of_day(arrival)
        if arrival_minute < departure_minute:
            arrival_minute += 24 * 60
        legs.append(
            {
                "id": f"flight-{flight}",
                "route": route,
                "line": _clean_text(value("Line"), "Line", row_number),
                "fleet": _clean_text(value("Fleet"), "Fleet", row_number),
                "day": _integer(value("Day"), "Day", row_number),
                "sequenceWithinRoute": route_sequences[route],
                "pairing": _integer(value("Pairing"), "Pairing", row_number),
                "flight": flight,
                "origin": _clean_text(value("Origin"), "Origin", row_number).upper(),
                "destination": _clean_text(value("Dest"), "Dest", row_number).upper(),
                "departure": departure,
                "arrival": arrival,
                "departureMinute": departure_minute,
                "arrivalMinute": arrival_minute,
            }
        )
    if not legs:
        raise ValueError("Routings sheet contains no flights")
    return legs


def _cyclic_distance(first: int | float, second: int | float) -> float:
    return min((first - second) % 1440, (second - first) % 1440)


def apply_gate_timing_baseline(
    legs: list[dict[str, Any]], gate_baseline: dict[str, Any]
) -> None:
    """Restore operational minute precision lost in the workbook export.

    The legacy workbook rounds times to ``HH:MM`` while the published gate
    artifact retains the numeric values used by the assignment engine. This
    one-time migration adapter matches each leg to its arrival and departure
    claim by route, city, adjacent city, and rounded clock value. Future native
    builds will write the operational minute fields directly.
    """
    claims_by_city = {
        city["code"]: city.get("claims", []) for city in gate_baseline["cities"]
    }
    for leg in legs:
        route = str(leg["route"])
        departure_clock = _minute_of_day(leg["departure"])
        arrival_clock = _minute_of_day(leg["arrival"])

        arrival_candidates = [
            claim["start"]
            for claim in claims_by_city.get(leg["destination"], [])
            if str(claim["label"]).split(" -> ")[0] == route
            and claim.get("arrivalCity") == leg["origin"]
        ]
        departure_candidates = [
            claim["end"]
            for claim in claims_by_city.get(leg["origin"], [])
            if str(claim["label"]).split(" -> ")[-1] == route
            and claim.get("departureCity") == leg["destination"]
        ]
        if not arrival_candidates or not departure_candidates:
            raise ValueError(
                "Gate timing baseline has no claim match for "
                f"flight {leg['flight']} ({leg['origin']}-{leg['destination']})"
            )

        arrival = min(
            arrival_candidates,
            key=lambda candidate: (_cyclic_distance(candidate, arrival_clock), candidate),
        )
        departure = min(
            departure_candidates,
            key=lambda candidate: (
                _cyclic_distance(candidate, departure_clock),
                candidate,
            ),
        )
        arrival_error = _cyclic_distance(arrival, arrival_clock)
        departure_error = _cyclic_distance(departure, departure_clock)
        if arrival_error > 1 or departure_error > 1:
            raise ValueError(
                "Gate timing baseline differs from workbook by more than one minute "
                f"for flight {leg['flight']} (departure {departure_error:.3f}, "
                f"arrival {arrival_error:.3f})"
            )

        operational_departure = departure % 1440
        operational_arrival = arrival % 1440
        if operational_arrival < operational_departure:
            operational_arrival += 1440
        leg["departureMinute"] = operational_departure
        leg["arrivalMinute"] = operational_arrival


def import_canonical_schedule(
    *,
    workbook_path: Path,
    city_information_path: Path,
    schedule: dict[str, Any],
    connection_window: dict[str, int],
    gate_plan: dict[str, Any] | None = None,
    gate_baseline_path: Path | None = None,
) -> dict[str, Any]:
    cities = load_city_information(city_information_path)
    legs = load_routings(workbook_path)
    if gate_baseline_path is not None:
        import json

        gate_baseline = json.loads(gate_baseline_path.read_text(encoding="utf-8"))
        apply_gate_timing_baseline(legs, gate_baseline)
    provenance: dict[str, Any] = {
        "workbook": {
            "filename": workbook_path.name,
            "sha256": sha256_file(workbook_path),
            "sheet": ROUTINGS_SHEET,
        },
        "cityInformation": {
            "filename": city_information_path.name,
            "sha256": sha256_file(city_information_path),
        },
    }
    if gate_baseline_path is not None:
        provenance["gateTimingBaseline"] = {
            "filename": gate_baseline_path.name,
            "sha256": sha256_file(gate_baseline_path),
        }
    return {
        "schemaVersion": "1.0.0",
        "schedule": {
            "id": schedule["id"],
            "number": int(schedule["number"]),
            "version": str(schedule["version"]),
            "label": schedule["label"],
            "status": schedule["status"],
            "connectionWindowMinutes": {
                "minimum": int(connection_window["minimum"]),
                "maximum": int(connection_window["maximum"]),
            },
        },
        "provenance": provenance,
        "gatePlan": gate_plan
        or {
            "label": schedule["label"],
            "forcedStandSplits": {},
        },
        "cities": cities,
        "legs": legs,
    }
